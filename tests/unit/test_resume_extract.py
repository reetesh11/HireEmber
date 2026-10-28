import pytest
from docx import Document
from pypdf import PdfWriter

from hireember.resume_extract import extract_docx_text, extract_pdf_text, propose_profile


def test_blank_page_pdf_gives_empty(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "blank.pdf"
    with path.open("wb") as fh:
        writer.write(fh)
    assert extract_pdf_text(path) == ""


def test_encrypted_pdf_gives_empty(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    path = tmp_path / "enc.pdf"
    with path.open("wb") as fh:
        writer.write(fh)
    assert extract_pdf_text(path) == ""


def test_corrupt_and_missing_give_empty(tmp_path):
    bad = tmp_path / "x.pdf"
    bad.write_bytes(b"garbage")
    assert extract_pdf_text(bad) == ""
    assert extract_pdf_text(tmp_path / "none.pdf") == ""


def test_year_range_is_not_a_phone():
    assert "phone" not in propose_profile("Jane Doe\nEngineer 2019 - 2023\n")


def test_empty_text():
    assert propose_profile("") == {}


def test_single_phone_and_linkedin_as_written():
    got = propose_profile("Jane Doe\n+1 (416) 555-0142 | www.linkedin.com/in/jane\n")
    assert got["phone"] == "+1 (416) 555-0142"
    assert got["linkedin_url"] == "www.linkedin.com/in/jane"


def test_docx_paragraphs_and_table_in_order(tmp_path):
    doc = Document()
    doc.add_paragraph("Intro")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Acme"
    table.rows[0].cells[1].text = "Engineer"
    doc.add_paragraph("Closing")
    path = tmp_path / "cv.docx"
    doc.save(str(path))
    assert extract_docx_text(path).splitlines() == ["Intro", "Acme\tEngineer", "Closing"]


def test_corrupt_docx_raises_value_error(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"nope")
    with pytest.raises(ValueError, match="could not be read"):
        extract_docx_text(path)
